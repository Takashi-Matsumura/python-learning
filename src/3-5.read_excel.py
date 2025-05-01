# 3-5.read_excel.py

from openpyxl import load_workbook

wb = load_workbook("output/sales.xlsx")
ws = wb.active

for row in ws.iter_rows(values_only=True):
    print(row)
